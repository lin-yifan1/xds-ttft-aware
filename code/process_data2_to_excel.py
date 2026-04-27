from __future__ import annotations

import argparse
import hashlib
import math
from pathlib import Path
import re
import sqlite3

import pandas as pd
from openpyxl import Workbook

from config import (
    EXCEL_SHEET_MAX_LEN,
    PROCESSED_WORKBOOK_PATH,
    RAW_DATA_DIR,
    SHEET_NAME_FALLBACK_LEFT_BUDGET,
    SHEET_NAME_LEFT_BUDGET,
    SHEET_NAME_MIN_RIGHT_BUDGET,
    SHEET_NAME_SEPARATOR,
)

REQUIRED_COLUMNS = [
    "infer_service_id",
    "service_name",
    "domain_id",
    "rpm",
    "tpm",
    "ttft_avg",
    "tpot_avg",
    "prompt_tokens",
    "completion_tokens",
    "collect_time_std",
]

OUTPUT_COLUMNS = [
    "domain_id",
    "rpm",
    "tpm",
    "ttft_avg",
    "tpot_avg",
    "prompt_tokens",
    "completion_tokens",
    "collect_time_std",
]

NUMERIC_COLUMNS = [
    "rpm",
    "tpm",
    "ttft_avg",
    "tpot_avg",
    "prompt_tokens",
    "completion_tokens",
]

SQL_TABLE = "processed_rows"
SQL_COLUMNS = [
    "infer_service_id",
    "service_name",
    "domain_id",
    "rpm",
    "tpm",
    "ttft_avg",
    "tpot_avg",
    "prompt_tokens",
    "completion_tokens",
    "collect_time_std",
]
CSV_ENCODINGS = ("utf-8", "utf-8-sig", "gb18030", "gbk")
DEFAULT_CHUNK_ROWS = 100_000
EXCEL_MAX_ROWS = 1_048_576
EXCEL_MAX_DATA_ROWS = EXCEL_MAX_ROWS - 1

_INVALID_SHEET_CHARS = re.compile(r"[\[\]\:\*\?\/\\]")


def sanitize_excel_sheet_name(name: str, max_len: int = EXCEL_SHEET_MAX_LEN) -> str:
    s = _INVALID_SHEET_CHARS.sub("_", str(name).strip())
    s = s.strip("'") or "sheet"
    if max_len < 1:
        max_len = 1
    return s[:max_len]


def reserve_sheet_name(name: str, used: set[str]) -> str:
    candidate = sanitize_excel_sheet_name(name, EXCEL_SHEET_MAX_LEN)
    if candidate not in used:
        used.add(candidate)
        return candidate

    for i in range(2, 10000):
        suffix = f"_{i}"
        candidate = sanitize_excel_sheet_name(
            f"{name[: EXCEL_SHEET_MAX_LEN - len(suffix)]}{suffix}",
            EXCEL_SHEET_MAX_LEN,
        )
        if candidate not in used:
            used.add(candidate)
            return candidate

    raise RuntimeError("Unable to generate a unique sheet name.")


def build_sheet_name(infer_service_id: str, service_name: str, used: set[str]) -> str:
    sep = SHEET_NAME_SEPARATOR
    sep_len = len(sep)

    def readable_name() -> str:
        left_budget = SHEET_NAME_LEFT_BUDGET
        right_budget = EXCEL_SHEET_MAX_LEN - sep_len - left_budget
        if right_budget < SHEET_NAME_MIN_RIGHT_BUDGET:
            left_budget = SHEET_NAME_FALLBACK_LEFT_BUDGET
            right_budget = EXCEL_SHEET_MAX_LEN - sep_len - left_budget
        left = sanitize_excel_sheet_name(infer_service_id, left_budget)
        right = sanitize_excel_sheet_name(service_name, right_budget)
        return sanitize_excel_sheet_name(f"{left}{sep}{right}", EXCEL_SHEET_MAX_LEN)

    name = readable_name()
    if name not in used:
        used.add(name)
        return name

    digest = hashlib.sha256(f"{infer_service_id}\0{service_name}".encode("utf-8")).hexdigest()[:8]
    fallback = sanitize_excel_sheet_name(f"{sanitize_excel_sheet_name(infer_service_id, 10)}{sep}h{digest}")
    if fallback not in used:
        used.add(fallback)
        return fallback

    for i in range(2, 10000):
        suffix = f"_{i}"
        candidate = sanitize_excel_sheet_name(name[: EXCEL_SHEET_MAX_LEN - len(suffix)] + suffix, EXCEL_SHEET_MAX_LEN)
        if candidate not in used:
            used.add(candidate)
            return candidate

    raise RuntimeError("Unable to generate a unique sheet name.")


def build_part_sheet_name(base_name: str, part_idx: int, used: set[str]) -> str:
    suffix = f"_p{part_idx}"
    name = f"{base_name[: EXCEL_SHEET_MAX_LEN - len(suffix)]}{suffix}"
    return reserve_sheet_name(name, used)


def cleanup_sqlite_files(db_path: Path) -> None:
    for path in (
        db_path,
        Path(f"{db_path}-wal"),
        Path(f"{db_path}-shm"),
        Path(f"{db_path}-journal"),
    ):
        if path.exists():
            path.unlink()


def open_temp_database(db_path: Path) -> sqlite3.Connection:
    db_path.parent.mkdir(parents=True, exist_ok=True)
    cleanup_sqlite_files(db_path)
    conn = sqlite3.connect(db_path)
    conn.execute("PRAGMA journal_mode=OFF")
    conn.execute("PRAGMA synchronous=OFF")
    conn.execute("PRAGMA temp_store=FILE")
    conn.execute("PRAGMA cache_size=-65536")
    conn.execute("PRAGMA locking_mode=EXCLUSIVE")
    conn.execute(
        f"""
        CREATE TABLE {SQL_TABLE} (
            infer_service_id TEXT NOT NULL,
            service_name TEXT NOT NULL,
            domain_id TEXT NOT NULL,
            rpm REAL NOT NULL,
            tpm REAL NOT NULL,
            ttft_avg REAL NOT NULL,
            tpot_avg REAL NOT NULL,
            prompt_tokens REAL NOT NULL,
            completion_tokens REAL NOT NULL,
            collect_time_std TEXT NOT NULL
        )
        """
    )
    return conn


def parse_collect_time_series(values: pd.Series) -> pd.Series:
    text = values.fillna("").astype(str).str.strip().str.replace(r"\s+", " ", regex=True)
    parsed = pd.to_datetime(text, format="%Y-%m-%d %H:%M:%S", errors="coerce")

    for fmt in ("%Y/%m/%d %H:%M:%S", "%Y-%m-%d %H:%M", "%Y/%m/%d %H:%M"):
        mask = parsed.isna() & text.ne("")
        if not mask.any():
            return parsed
        parsed.loc[mask] = pd.to_datetime(text.loc[mask], format=fmt, errors="coerce")

    mask = parsed.isna() & text.ne("")
    if mask.any():
        parsed.loc[mask] = pd.to_datetime(text.loc[mask], errors="coerce")
    return parsed


def clean_text_series(values: pd.Series) -> pd.Series:
    return values.fillna("").astype(str).str.strip()


def normalize_chunk(chunk: pd.DataFrame) -> tuple[pd.DataFrame, int, int]:
    missing = [c for c in REQUIRED_COLUMNS if c not in chunk.columns]
    if missing:
        raise ValueError(f"Missing required columns: {missing}")

    infer_service_id = clean_text_series(chunk["infer_service_id"])
    service_name = clean_text_series(chunk["service_name"])
    domain_id = clean_text_series(chunk["domain_id"])
    parsed_time = parse_collect_time_series(chunk["collect_time_std"])

    valid_time = parsed_time.notna()
    valid_keys = infer_service_id.ne("") & service_name.ne("") & domain_id.ne("")
    keep = valid_time & valid_keys
    kept = int(keep.sum())
    dropped_time = int((~valid_time).sum())
    dropped_keys = int((valid_time & ~valid_keys).sum())

    if kept == 0:
        return pd.DataFrame(columns=SQL_COLUMNS), dropped_time, dropped_keys

    idx = keep[keep].index
    normalized = pd.DataFrame(
        {
            "infer_service_id": infer_service_id.loc[idx].to_numpy(),
            "service_name": service_name.loc[idx].to_numpy(),
            "domain_id": domain_id.loc[idx].to_numpy(),
            "collect_time_std": parsed_time.loc[idx].dt.strftime("%Y-%m-%d %H:%M:%S").to_numpy(),
        }
    )

    for col in NUMERIC_COLUMNS:
        normalized[col] = (
            pd.to_numeric(chunk.loc[idx, col], errors="coerce")
            .fillna(0.0)
            .astype("float64")
            .to_numpy()
        )

    return normalized[SQL_COLUMNS], dropped_time, dropped_keys


def insert_rows(conn: sqlite3.Connection, frame: pd.DataFrame) -> int:
    if frame.empty:
        return 0

    placeholders = ", ".join("?" for _ in SQL_COLUMNS)
    columns = ", ".join(SQL_COLUMNS)
    conn.executemany(
        f"INSERT INTO {SQL_TABLE} ({columns}) VALUES ({placeholders})",
        frame[SQL_COLUMNS].itertuples(index=False, name=None),
    )
    return len(frame)


def read_csv_chunks(path: Path, encoding: str, chunk_rows: int):
    return pd.read_csv(
        path,
        encoding=encoding,
        usecols=REQUIRED_COLUMNS,
        chunksize=chunk_rows,
        low_memory=False,
    )


def load_csv_file(conn: sqlite3.Connection, idx: int, total: int, path: Path, chunk_rows: int) -> tuple[int, int, int]:
    print(f"[progress] reading file {idx}/{total}: {path.name}")
    last_error: Exception | None = None

    for encoding in CSV_ENCODINGS:
        read_rows = 0
        kept_rows = 0
        dropped_time = 0
        dropped_keys = 0
        chunk_no = 0
        conn.execute("SAVEPOINT load_file")
        try:
            for chunk in read_csv_chunks(path, encoding, chunk_rows):
                chunk_no += 1
                read_rows += len(chunk)
                normalized, bad_time, bad_keys = normalize_chunk(chunk)
                kept_rows += insert_rows(conn, normalized)
                dropped_time += bad_time
                dropped_keys += bad_keys

                if chunk_no == 1 or chunk_no % 10 == 0:
                    print(
                        f"       chunks={chunk_no} read_rows={read_rows} "
                        f"kept_rows={kept_rows}"
                    )

            conn.execute("RELEASE load_file")
            print(
                f"       rows={read_rows} kept={kept_rows} "
                f"dropped_time={dropped_time} dropped_empty_key={dropped_keys} "
                f"encoding={encoding}"
            )
            return read_rows, kept_rows, dropped_time + dropped_keys
        except Exception as exc:
            conn.execute("ROLLBACK TO load_file")
            conn.execute("RELEASE load_file")
            if isinstance(exc, ValueError) and (
                "Missing required columns" in str(exc) or "Usecols do not match" in str(exc)
            ):
                raise
            last_error = exc

    raise RuntimeError(f"Failed to read csv: {path}") from last_error


def load_csv_dir(conn: sqlite3.Connection, input_dir: Path, chunk_rows: int, workers: int) -> int:
    files = sorted(input_dir.glob("*.csv"))
    if not files:
        raise FileNotFoundError(f"No csv files found under: {input_dir}")

    if workers not in (0, 1):
        print("[info] streaming mode keeps one CSV reader active at a time; --workers is accepted for compatibility.")
    print(f"[progress] loading csv files from {input_dir} total_files={len(files)} chunk_rows={chunk_rows}")

    total_read = 0
    total_kept = 0
    total_dropped = 0
    for idx, path in enumerate(files, start=1):
        read_rows, kept_rows, dropped_rows = load_csv_file(conn, idx, len(files), path, chunk_rows)
        total_read += read_rows
        total_kept += kept_rows
        total_dropped += dropped_rows

    print(f"[load] read rows={total_read} kept rows={total_kept} dropped rows={total_dropped}")
    if total_kept == 0:
        raise ValueError("No valid rows remained after filtering.")
    return total_kept


def create_sort_index(conn: sqlite3.Connection) -> None:
    print("[progress] creating sort index for grouped Excel export")
    conn.execute(
        f"""
        CREATE INDEX idx_processed_group_order
        ON {SQL_TABLE} (infer_service_id, service_name, domain_id, collect_time_std)
        """
    )
    conn.commit()


def fetch_sheet_groups(conn: sqlite3.Connection) -> list[tuple[str, str, int]]:
    groups = conn.execute(
        f"""
        SELECT infer_service_id, service_name, COUNT(*) AS row_count
        FROM {SQL_TABLE}
        GROUP BY infer_service_id, service_name
        ORDER BY infer_service_id, service_name
        """
    ).fetchall()
    print(f"[progress] discovered sheet groups={len(groups)}")
    return [(str(infer), str(service), int(row_count)) for infer, service, row_count in groups]


def append_group_rows(
    conn: sqlite3.Connection,
    workbook: Workbook,
    used_sheet_names: set[str],
    idx: int,
    total: int,
    infer_service_id: str,
    service_name: str,
    row_count: int,
) -> int:
    base_sheet_name = build_sheet_name(infer_service_id, service_name, used_sheet_names)
    part_count = max(1, math.ceil(row_count / EXCEL_MAX_DATA_ROWS))
    if part_count > 1:
        print(
            f"[warn] group {idx}/{total} has {row_count} rows; "
            f"splitting across {part_count} sheets"
        )

    print(
        f"[progress] writing group {idx}/{total} "
        f"infer_service_id={infer_service_id} service_name={service_name} rows={row_count}"
    )
    cursor = conn.execute(
        f"""
        SELECT domain_id, rpm, tpm, ttft_avg, tpot_avg, prompt_tokens, completion_tokens, collect_time_std
        FROM {SQL_TABLE}
        WHERE infer_service_id = ? AND service_name = ?
        ORDER BY domain_id, collect_time_std
        """,
        (infer_service_id, service_name),
    )

    sheet_count = 0
    rows_in_sheet = EXCEL_MAX_DATA_ROWS
    ws = None
    for row in cursor:
        if rows_in_sheet >= EXCEL_MAX_DATA_ROWS:
            sheet_count += 1
            sheet_name = (
                base_sheet_name
                if sheet_count == 1
                else build_part_sheet_name(base_sheet_name, sheet_count, used_sheet_names)
            )
            ws = workbook.create_sheet(sheet_name)
            ws.append(OUTPUT_COLUMNS)
            rows_in_sheet = 0
            print(f"[sheet] {sheet_name} start part={sheet_count}/{part_count}")

        ws.append(row)
        rows_in_sheet += 1

    print(f"[sheet] {base_sheet_name} rows={row_count} sheets={sheet_count}")
    return sheet_count


def write_excel_from_db(conn: sqlite3.Connection, output_path: Path) -> Path:
    output_path.parent.mkdir(parents=True, exist_ok=True)
    temp_path = output_path.with_name(f"{output_path.stem}.tmp{output_path.suffix}")
    if temp_path.exists():
        temp_path.unlink()

    groups = fetch_sheet_groups(conn)
    if not groups:
        raise ValueError("No valid rows remained after filtering.")

    print(f"[progress] writing workbook sheets={len(groups)} output={output_path}")
    workbook = Workbook(write_only=True)
    used_sheet_names: set[str] = set()
    written_sheets = 0
    for idx, (infer_service_id, service_name, row_count) in enumerate(groups, start=1):
        written_sheets += append_group_rows(
            conn,
            workbook,
            used_sheet_names,
            idx,
            len(groups),
            infer_service_id,
            service_name,
            row_count,
        )

    workbook.save(temp_path)
    temp_path.replace(output_path)
    print(f"[ok] wrote {output_path} sheets={written_sheets}")
    return output_path


def default_temp_db_path(output_path: Path) -> Path:
    return output_path.with_name(f"{output_path.stem}.work.sqlite")


def parse_args() -> argparse.Namespace:
    ap = argparse.ArgumentParser(description="Filter data2 csv files and write grouped excel sheets.")
    ap.add_argument("--input-dir", type=Path, default=RAW_DATA_DIR, help="Directory containing source csv files.")
    ap.add_argument("--output", type=Path, default=PROCESSED_WORKBOOK_PATH, help="Output xlsx path.")
    ap.add_argument(
        "--workers",
        type=int,
        default=0,
        help="Accepted for compatibility. The streaming pipeline reads one CSV file at a time to keep memory stable.",
    )
    ap.add_argument(
        "--chunk-rows",
        type=int,
        default=DEFAULT_CHUNK_ROWS,
        help="CSV rows to process per chunk.",
    )
    ap.add_argument(
        "--temp-db",
        type=Path,
        default=None,
        help="Optional path for the temporary SQLite database.",
    )
    ap.add_argument(
        "--keep-temp-db",
        action="store_true",
        help="Keep the temporary SQLite database after a successful run for debugging.",
    )
    return ap.parse_args()


def main() -> int:
    args = parse_args()
    if args.chunk_rows < 1:
        raise ValueError("--chunk-rows must be >= 1")

    temp_db_path = args.temp_db or default_temp_db_path(args.output)
    print(f"[progress] using temporary sqlite database {temp_db_path}")
    conn = open_temp_database(temp_db_path)
    try:
        load_csv_dir(conn, args.input_dir, args.chunk_rows, args.workers)
        create_sort_index(conn)
        write_excel_from_db(conn, args.output)
    finally:
        conn.close()
        if not args.keep_temp_db:
            cleanup_sqlite_files(temp_db_path)

    return 0


if __name__ == "__main__":
    raise SystemExit(main())
